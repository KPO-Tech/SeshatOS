"""A light, model-free first tier of text extraction for connectors.

Most business documents have a real text layer or are plain Office files, and they do not need a
layout model. This tier reads them with small libraries (PDFium, python-docx, python-pptx) and leaves
scanned PDFs and anything unusual to the heavy tier (the Docling pool), which is slow and memory hungry.

Safety rules, because these bytes come from other people's drives:
- Office files are zip containers, so size, ratio and compression type are checked before parsing.
- PDFium can hang or abort on a malformed PDF, so it runs in a worker process under a deadline. A
  timeout or crash replaces the worker pool instead of taking the service down.
- An encrypted PDF is reported, not silently emptied.
"""

from __future__ import annotations

import asyncio
import io
import os
import zipfile
from concurrent.futures import ProcessPoolExecutor
from concurrent.futures.process import BrokenProcessPool
from typing import Any, Callable

from seshat_intelligence.connectors.extraction import ExtractionError, Extractor

MAX_ZIP_MEMBER_BYTES = 200 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 1024 * 1024 * 1024
MAX_ZIP_RATIO = 200
_RATIO_CHECK_MIN_BYTES = 1024 * 1024
# zipfile decompresses bzip2 and lzma without an output bound, so only these are allowed.
_ALLOWED_COMPRESSION = {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}

# A PDF whose pages average fewer characters than this has no usable text layer (a scan).
MIN_CHARS_PER_PAGE = 20
PDF_TIMEOUT_SECONDS = 60.0


def check_zip_limits(data: bytes) -> None:
    """Refuse zip bombs: oversized members, oversized totals, absurd compression ratios."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractionError("not a valid Office file") from exc
    total = 0
    for info in archive.infolist():
        if info.is_dir():
            continue
        if info.compress_type not in _ALLOWED_COMPRESSION:
            raise ExtractionError("archive uses an unsupported compression method")
        if info.file_size > MAX_ZIP_MEMBER_BYTES:
            raise ExtractionError("archive member is too large")
        total += info.file_size
        if total > MAX_ZIP_TOTAL_BYTES:
            raise ExtractionError("archive is too large once decompressed")
        if info.file_size > _RATIO_CHECK_MIN_BYTES and info.file_size > max(1, info.compress_size) * MAX_ZIP_RATIO:
            raise ExtractionError("archive is compressed suspiciously well")


# -- workers (module level so they can be pickled into a process pool) --------------------------


def _pdf_text(data: bytes) -> tuple[str, int]:
    """Text of every page and the page count. Raises ExtractionError for an encrypted file."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    password = None
    if reader.is_encrypted:
        # Owner-password-only files (copy restrictions, no open password) decrypt with "".
        if reader.decrypt("") == 0:
            raise ExtractionError("the PDF is encrypted")
        password = ""
    try:
        import pypdfium2 as pdfium

        pdf = pdfium.PdfDocument(data, password=password)
        try:
            pages: list[str] = []
            for page in pdf:
                text_page = page.get_textpage()
                try:
                    pages.append(text_page.get_text_range())
                finally:
                    text_page.close()
                    page.close()
            return "\n\n".join(pages), len(pages)
        finally:
            pdf.close()
    except Exception:  # noqa: BLE001 - PDFium failed on this file, pypdf is the fallback
        texts = [(page.extract_text() or "") for page in reader.pages]
        return "\n\n".join(texts), len(texts)


def _docx_text(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                text = "#" * min(int(style[8:]), 6) + " " + text
            parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        if lines:
            slides.append(f"[Slide {number}]\n" + "\n".join(lines))
    return "\n\n".join(slides)


def _sleep_then_return(seconds: float) -> str:  # used by tests to exercise the deadline
    import time

    time.sleep(seconds)
    return "late"


class LightPool:
    """A small process pool with a deadline. A worker that hangs or crashes is discarded with its
    pool, and the next call gets a fresh one."""

    def __init__(self, max_workers: int = 2, timeout: float = PDF_TIMEOUT_SECONDS) -> None:
        self._max_workers = max_workers
        self._timeout = timeout
        self._executor = self._new_executor()

    def _new_executor(self) -> ProcessPoolExecutor:
        return ProcessPoolExecutor(max_workers=self._max_workers)

    def _replace(self) -> None:
        # The worker handles must be read before shutdown, which clears them. A worker stuck in a
        # native call ignores shutdown, so it is terminated outright.
        processes = list((getattr(self._executor, "_processes", None) or {}).values())
        self._executor.shutdown(wait=False, cancel_futures=True)
        for process in processes:
            try:
                process.terminate()
            except Exception:  # noqa: BLE001 - already gone
                pass
        self._executor = self._new_executor()

    async def run(self, fn: Callable[..., Any], *args: Any, timeout: float | None = None) -> Any:
        loop = asyncio.get_running_loop()
        try:
            return await asyncio.wait_for(loop.run_in_executor(self._executor, fn, *args), timeout or self._timeout)
        except asyncio.TimeoutError as exc:
            self._replace()
            raise ExtractionError("extraction took too long") from exc
        except BrokenProcessPool as exc:
            self._replace()
            raise ExtractionError("the extraction worker crashed") from exc

    def shutdown(self) -> None:
        self._executor.shutdown(wait=False, cancel_futures=True)


class LightExtractor:
    """Extract text from PDF, DOCX and PPTX without a model. Returns None when this tier cannot
    handle a file (another format, or a PDF with no text layer) so the caller can fall back."""

    def __init__(self, pool: LightPool | None = None) -> None:
        self._pool = pool or LightPool()

    async def __call__(self, data: bytes, mime_type: str, name: str) -> str | None:
        extension = os.path.splitext(name)[1].lower()
        if extension == ".pdf":
            text, pages = await self._pool.run(_pdf_text, data)
            if pages == 0 or len(text.strip()) / pages < MIN_CHARS_PER_PAGE:
                return None  # a scan: leave it to the OCR-capable tier
            return text
        if extension == ".docx":
            check_zip_limits(data)
            return await self._pool.run(_docx_text, data)
        if extension == ".pptx":
            check_zip_limits(data)
            return await self._pool.run(_pptx_text, data)
        return None

    def shutdown(self) -> None:
        self._pool.shutdown()


def tiered_extractor(light: Extractor, heavy: Extractor | None) -> Extractor:
    """Try the light tier first and use the heavy one only when the light tier declines."""

    async def extract(data: bytes, mime_type: str, name: str) -> str | None:
        text = await light(data, mime_type, name)
        if text is not None and text.strip():
            return text
        if heavy is None:
            return text
        return await heavy(data, mime_type, name)

    return extract
