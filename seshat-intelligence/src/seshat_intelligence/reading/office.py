"""Native text for Office files (DOCX, PPTX, XLSX), no model involved.

These are zip containers from other people's drives, so size, ratio and compression type are checked
before anything is parsed. The functions here run inside the reading pool's worker processes.
"""

from __future__ import annotations

import io
import zipfile

from seshat_intelligence.reading.models import ReadError, UnsafeFile

MAX_ZIP_MEMBER_BYTES = 200 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 1024 * 1024 * 1024
MAX_ZIP_RATIO = 200
_RATIO_CHECK_MIN_BYTES = 1024 * 1024
# zipfile decompresses bzip2 and lzma without an output bound, so only these are allowed.
_ALLOWED_COMPRESSION = {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}

MAX_SHEET_ROWS = 5000
MAX_SHEET_COLUMNS = 50


def check_zip_limits(data: bytes) -> None:
    """Refuse zip bombs: oversized members, oversized totals, absurd compression ratios."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ReadError("not a valid Office file") from exc
    total = 0
    for info in archive.infolist():
        if info.is_dir():
            continue
        if info.compress_type not in _ALLOWED_COMPRESSION:
            raise UnsafeFile("archive uses an unsupported compression method")
        if info.file_size > MAX_ZIP_MEMBER_BYTES:
            raise UnsafeFile("archive member is too large")
        total += info.file_size
        if total > MAX_ZIP_TOTAL_BYTES:
            raise UnsafeFile("archive is too large once decompressed")
        if info.file_size > _RATIO_CHECK_MIN_BYTES and info.file_size > max(1, info.compress_size) * MAX_ZIP_RATIO:
            raise UnsafeFile("archive is compressed suspiciously well")


def docx_text(data: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            if style.startswith("Heading ") and style[8:].isdigit():
                text = "#" * min(int(style[8:]), 6) + " " + text
            parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def pptx_text(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        if lines:
            slides.append(f"[Slide {number}]\n" + "\n".join(lines))
    return "\n\n".join(slides)


def xlsx_text(data: bytes) -> str:
    """Each sheet as a heading and pipe-separated rows. Empty rows and trailing empty columns are
    dropped, and a sheet is capped so a huge export cannot flood the output."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= MAX_SHEET_ROWS:
                    rows.append("...")
                    break
                cells = ["" if value is None else str(value).strip().replace("\n", " ") for value in row[:MAX_SHEET_COLUMNS]]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"## {worksheet.title}\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def office_text(extension: str, data: bytes) -> str:
    """Native text for a supported Office extension, or raises ReadError on an unsafe file."""
    check_zip_limits(data)
    if extension == ".docx":
        return docx_text(data)
    if extension == ".pptx":
        return pptx_text(data)
    if extension == ".xlsx":
        return xlsx_text(data)
    raise ReadError(f"{extension} is not a native Office format")
