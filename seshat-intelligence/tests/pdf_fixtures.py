"""Generated PDFs and Office files for the reading tests."""

import io

from docx import Document as DocxDocument
from fpdf import FPDF
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

LONG_TEXT = "Plain paragraph with plenty of real words so that it counts as a genuine text layer."


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (400, 400), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


def make_pdf(kinds: list[str]) -> bytes:
    """One page per entry: text, scan (full-page image, no text), logo (tiny image plus text),
    empty (nothing), garbled (text full of cid placeholders)."""
    pdf = FPDF()
    for kind in kinds:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        if kind == "text":
            pdf.multi_cell(0, 8, LONG_TEXT)
        elif kind == "scan":
            pdf.image(io.BytesIO(_png()), x=0, y=0, w=210, h=297)
        elif kind == "logo":
            pdf.image(io.BytesIO(_png()), x=10, y=10, w=10, h=10)
            pdf.multi_cell(0, 8, LONG_TEXT)
        elif kind == "garbled":
            pdf.multi_cell(0, 8, "(cid:12)(cid:47)(cid:3) " * 6)
        elif kind == "empty":
            pass
        else:
            raise ValueError(kind)
    return bytes(pdf.output())


def encrypted_pdf(user_password: str, owner_password: str | None = None) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["text"]))))
    writer.encrypt(user_password=user_password, owner_password=owner_password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Refund policy", level=1)
    document.add_paragraph("Restocking fee is 10% for non-defective returns.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Fee"
    table.cell(1, 0).text, table.cell(1, 1).text = "Laptop", "10%"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Roadmap"
    first.placeholders[1].text = "Ship connectors"
    first.notes_slide.notes_text_frame.text = "Mention the pilot"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Risks"
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fees"
    sheet.append(["Item", "Fee"])
    sheet.append(["Laptop", "10%"])
    sheet.append([None, None])
    sheet.append(["Phone", "5%"])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
