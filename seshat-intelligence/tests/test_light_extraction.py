import io
import zipfile

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from seshat_intelligence.connectors.extraction import ExtractionError
from seshat_intelligence.connectors.light_extraction import (
    LightExtractor,
    LightPool,
    _sleep_then_return,
    check_zip_limits,
    tiered_extractor,
)


def make_pdf(pages: list[str]) -> bytes:
    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        if text:
            pdf.multi_cell(0, 8, text)
        else:
            pdf.rect(20, 20, 50, 50)  # drawing only: no text layer
    return bytes(pdf.output())


def make_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Refund policy", level=1)
    document.add_paragraph("Restocking fee is 10% for non-defective returns.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Fee"
    table.cell(1, 0).text, table.cell(1, 1).text = "Laptop", "10%"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Roadmap"
    first.placeholders[1].text = "Ship connectors"
    first.notes_slide.notes_text_frame.text = "Mention the pilot"
    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Risks"
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


@pytest.fixture
async def extractor():
    light = LightExtractor(LightPool(max_workers=1, timeout=30))
    yield light
    light.shutdown()


async def test_pdf_with_a_text_layer(extractor):
    pdf = make_pdf(["First page about refunds and restocking fees.", "Second page about shipping times."])
    text = await extractor(pdf, "application/pdf", "policy.pdf")
    assert "restocking fees" in text and "shipping times" in text


async def test_scanned_pdf_is_declined_so_the_heavy_tier_can_take_it(extractor):
    assert await extractor(make_pdf([""]), "application/pdf", "scan.pdf") is None


async def test_encrypted_pdf_is_reported(extractor):
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["Secret content for the owner only."]))))
    writer.encrypt(user_password="open-sesame")
    out = io.BytesIO()
    writer.write(out)
    with pytest.raises(ExtractionError, match="encrypted"):
        await extractor(out.getvalue(), "application/pdf", "locked.pdf")


async def test_owner_password_only_pdf_is_readable(extractor):
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["Readable although copy is restricted."]))))
    writer.encrypt(user_password="", owner_password="owner")
    out = io.BytesIO()
    writer.write(out)
    assert "Readable" in await extractor(out.getvalue(), "application/pdf", "restricted.pdf")


async def test_docx_keeps_headings_paragraphs_and_tables_in_order(extractor):
    text = await extractor(make_docx(), "", "policy.docx")
    assert text.index("# Refund policy") < text.index("Restocking fee") < text.index("Item | Fee") < text.index("Laptop | 10%")


async def test_pptx_text_tables_and_notes(extractor):
    text = await extractor(make_pptx(), "", "deck.pptx")
    assert "[Slide 1]" in text and "Roadmap" in text and "Ship connectors" in text and "Notes: Mention the pilot" in text
    assert "[Slide 2]" in text and "Risks" in text


async def test_other_formats_are_declined(extractor):
    assert await extractor(b"x", "application/msword", "old.doc") is None


def test_zip_bomb_is_refused():
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (20 * 1024 * 1024))
    with pytest.raises(ExtractionError, match="compressed suspiciously"):
        check_zip_limits(bomb.getvalue())


def test_unsupported_compression_is_refused():
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, "w", zipfile.ZIP_BZIP2) as archive:
        archive.writestr("a.xml", b"data")
    with pytest.raises(ExtractionError, match="compression"):
        check_zip_limits(packed.getvalue())


def test_a_normal_office_file_passes_the_zip_check():
    check_zip_limits(make_docx())


async def test_a_hung_worker_times_out_and_the_pool_recovers():
    pool = LightPool(max_workers=1, timeout=30)
    try:
        with pytest.raises(ExtractionError, match="too long"):
            await pool.run(_sleep_then_return, 5, timeout=0.3)
        assert await pool.run(_sleep_then_return, 0) == "late"
    finally:
        pool.shutdown()


async def test_tiered_extractor_uses_the_heavy_tier_only_when_the_light_one_declines():
    calls = []

    async def light(data, mime, name):
        return "light text" if name.endswith(".docx") else None

    async def heavy(data, mime, name):
        calls.append(name)
        return "heavy text"

    extract = tiered_extractor(light, heavy)
    assert await extract(b"", "", "a.docx") == "light text"
    assert await extract(b"", "", "scan.pdf") == "heavy text"
    assert calls == ["scan.pdf"]
    assert await tiered_extractor(light, None)(b"", "", "scan.pdf") is None
