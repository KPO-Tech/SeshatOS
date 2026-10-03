import io
import zipfile

import pypdfium2 as pdfium
import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from pdf_fixtures import LONG_TEXT, encrypted_pdf, make_docx, make_pdf, make_pptx, make_xlsx

from seshat_intelligence.providers.base import ConvertedDocument
from seshat_intelligence.reading.engines import EnginePolicy
from seshat_intelligence.reading.models import PageSource, ReadError
from seshat_intelligence.reading.office import check_zip_limits
from seshat_intelligence.reading.pool import LightPool, sleep_then_return
from seshat_intelligence.reading.quality import is_garbled_text
from seshat_intelligence.reading.router import ReadingRouter
from seshat_intelligence.reading.routes import router as reading_routes


class FakeEngines:
    """Engines that record every call. `outcomes` maps an engine to a callable (filename, pages) ->
    ConvertedDocument, or to an exception to raise."""

    def __init__(self, order=("docling",), outcomes=None, available=None):
        self.order = list(order)
        self.outcomes = outcomes or {}
        self.engines = list(available) if available is not None else list(order)
        self.calls = []  # (engine, filename, page_count)

    def order_for(self, extension):
        return list(self.order)

    def available(self):
        return list(self.engines)

    async def convert(self, engine, filename, data):
        pages = len(pdfium.PdfDocument(data)) if filename.endswith(".pdf") else 0
        self.calls.append((engine, filename, pages))
        outcome = self.outcomes.get(engine)
        if isinstance(outcome, Exception):
            raise outcome
        if outcome is not None:
            return outcome(filename, pages)
        return ConvertedDocument(status="success", markdown=f"[{engine}] {filename}")


@pytest.fixture
async def make_router():
    routers = []

    def build(engines=None, **kwargs):
        router = ReadingRouter(engines or FakeEngines(), pool=LightPool(max_workers=1, timeout=60), **kwargs)
        routers.append(router)
        return router

    yield build
    for router in routers:
        router.shutdown()


# -- text quality, mirroring the Go tests ---------------------------------------------------------


def test_real_prose_is_not_garbled():
    assert not is_garbled_text("Bienvenue dans la lecon portant sur le diagnostic financier fonctionnel.")


def test_empty_text_is_not_garbled():
    assert not is_garbled_text("") and not is_garbled_text("   \n\t")


def test_cid_placeholders_are_garbled():
    assert is_garbled_text("Report Title (cid:47)(cid:12) Summary")


def test_text_dominated_by_private_use_characters_is_garbled():
    assert is_garbled_text("".join(chr(0xE000 + i) for i in range(50)))


def test_one_stray_private_use_character_is_not_garbled():
    assert not is_garbled_text("This is a perfectly normal sentence with real words. " * 20 + chr(0xE000))


# -- engine policy ---------------------------------------------------------------------------------


def test_policy_auto_tries_docling_then_marker_for_pdfs():
    assert EnginePolicy(["docling", "marker"], "auto").order_for(".pdf") == ["docling", "marker"]
    assert EnginePolicy(["marker", "docling"], "auto").order_for(".pdf") == ["docling", "marker"]


def test_policy_explicit_engine_is_used_alone_and_must_be_enabled():
    assert EnginePolicy(["docling", "marker"], "marker").order_for(".pdf") == ["marker"]
    assert EnginePolicy(["docling"], "marker").order_for(".pdf") == []


def test_policy_sends_everything_that_is_not_a_pdf_to_docling():
    policy = EnginePolicy(["docling", "marker"], "marker")
    assert policy.order_for(".html") == ["docling"]
    assert EnginePolicy(["marker"], "auto").order_for(".docx") == []


# -- pdf routing -----------------------------------------------------------------------------------


async def test_a_text_pdf_never_reaches_an_engine(make_router):
    engines = FakeEngines()
    result = await make_router(engines).read("a.pdf", make_pdf(["text", "text"]))
    assert result.ok and result.source == "native" and result.engines_used == []
    assert result.markdown.count(LONG_TEXT) == 2
    assert engines.calls == []
    assert [p.source for p in result.pages] == [PageSource.NATIVE, PageSource.NATIVE]


async def test_only_pages_that_need_it_go_to_the_engine(make_router):
    engines = FakeEngines()
    result = await make_router(engines).read("a.pdf", make_pdf(["text", "scan", "text"]))
    assert result.ok and result.source == "mixed"
    assert [(c[0], c[2]) for c in engines.calls] == [("docling", 1)]
    assert result.markdown.index(LONG_TEXT) < result.markdown.index("[docling]") < result.markdown.rindex(LONG_TEXT)
    assert [(p.page, p.source.value, p.reason) for p in result.pages] == [(1, "native", ""), (2, "engine", "image"), (3, "native", "")]
    assert result.engine_page_count == 1


async def test_a_logo_does_not_make_a_page_need_an_engine(make_router):
    engines = FakeEngines()
    result = await make_router(engines).read("a.pdf", make_pdf(["logo", "logo"]))
    assert result.source == "native" and engines.calls == []
    assert all(p.has_image is False for p in result.pages)


async def test_consecutive_heavy_pages_are_sent_in_one_call(make_router):
    engines = FakeEngines()
    result = await make_router(engines).read("a.pdf", make_pdf(["text", "scan", "empty", "text", "scan"]))
    assert [(c[1], c[2]) for c in engines.calls] == [("pages-2-3.pdf", 2), ("pages-5-5.pdf", 1)]
    assert [p.reason for p in result.pages if p.source == PageSource.ENGINE] == ["image", "sparse", "image"]


async def test_garbled_text_layer_goes_to_the_engine(make_router):
    engines = FakeEngines()
    result = await make_router(engines).read("a.pdf", make_pdf(["garbled"]))
    assert result.ok and result.source == "engine"
    assert result.pages[0].reason == "garbled"
    assert [c[2] for c in engines.calls] == [1]  # the single page is the whole document: sent as is


async def test_an_engine_answer_that_is_garbled_gets_a_second_opinion(make_router):
    def garbled(name, pages):
        return ConvertedDocument(status="success", markdown="(cid:1)(cid:2) broken")

    engines = FakeEngines(order=("docling", "marker"), outcomes={"docling": garbled})
    result = await make_router(engines).read("a.pdf", make_pdf(["text", "scan"]))
    assert result.ok and result.engines_used == ["marker"]
    assert [c[0] for c in engines.calls] == ["docling", "marker"]


async def test_a_crashing_engine_falls_through_to_the_next(make_router):
    engines = FakeEngines(order=("docling", "marker"), outcomes={"docling": RuntimeError("worker died")})
    result = await make_router(engines).read("a.pdf", make_pdf(["scan"]))
    assert result.ok and result.engines_used == ["marker"]


async def test_one_unreadable_page_discards_partial_results_and_the_whole_document_is_retried(make_router):
    def only_whole_documents(filename, pages):
        if filename.startswith("pages-"):
            return ConvertedDocument(status="failure", markdown="", errors=["no"])
        return ConvertedDocument(status="success", markdown="whole document text")

    engines = FakeEngines(outcomes={"docling": only_whole_documents})
    result = await make_router(engines).read("a.pdf", make_pdf(["text", "scan", "text"]))
    assert result.ok and result.markdown == "whole document text" and result.source == "engine"
    assert [c[1] for c in engines.calls] == ["pages-2-2.pdf", "a.pdf"]


async def test_when_even_the_whole_document_fails_the_result_is_not_ok(make_router):
    engines = FakeEngines(outcomes={"docling": lambda name, pages: ConvertedDocument(status="failure", markdown="")})
    result = await make_router(engines).read("a.pdf", make_pdf(["text", "scan"]))
    assert not result.ok and result.markdown == ""
    assert "whole-document fallback failed" in result.reason


async def test_pages_that_need_an_engine_with_none_available_are_not_silently_dropped(make_router):
    result = await make_router(FakeEngines(order=())).read("a.pdf", make_pdf(["text", "scan"]))
    assert not result.ok and "1 page(s) need a conversion engine" in result.reason
    assert (await make_router(FakeEngines(order=())).read("a.pdf", make_pdf(["text"]))).ok  # text-only needs none


async def test_whole_mode_sends_the_document_to_the_engines(make_router):
    engines = FakeEngines()
    result = await make_router(engines, pdf_mode="whole").read("a.pdf", make_pdf(["text", "text"]))
    assert result.source == "engine" and [(c[1], c[2]) for c in engines.calls] == [("a.pdf", 2)]
    per_request = await make_router(FakeEngines()).read("a.pdf", make_pdf(["text"]), pdf_mode="whole")
    assert per_request.source == "engine"


async def test_encrypted_pdf_is_reported_and_owner_password_pdf_is_read(make_router):
    router = make_router()
    locked = await router.read("a.pdf", encrypted_pdf("open-sesame"))
    assert not locked.ok and "encrypted" in locked.reason
    restricted = await router.read("a.pdf", encrypted_pdf("", "owner"))
    assert restricted.ok and LONG_TEXT in restricted.markdown


async def test_a_corrupt_pdf_is_reported(make_router):
    result = await make_router().read("a.pdf", b"%PDF-1.4 not really")
    assert not result.ok


# -- reading modes ---------------------------------------------------------------------------------


async def test_docling_mode_sends_every_file_straight_to_docling(make_router):
    engines = FakeEngines(order=("docling", "marker"), available=["docling", "marker"])
    router = make_router(engines, mode="docling")
    await router.read("a.pdf", make_pdf(["text", "text"]))
    await router.read("p.docx", make_docx())
    await router.read("page.html", b"<p>x</p>")
    assert engines.calls == [("docling", "a.pdf", 2), ("docling", "p.docx", 0), ("docling", "page.html", 0)]
    assert (await router.read("n.md", b"# note")).source == "native"  # plain text never needs an engine


async def test_docling_mode_reports_when_docling_is_not_available(make_router):
    result = await make_router(FakeEngines(order=("marker",), available=["marker"]), mode="docling").read("a.pdf", make_pdf(["text"]))
    assert not result.ok and "docling is not available" in result.reason


async def test_marker_mode_sends_pdfs_to_marker_and_reads_other_formats_with_the_custom_reader(make_router):
    engines = FakeEngines(order=("docling", "marker"), available=["docling", "marker"])
    router = make_router(engines, mode="marker")
    pdf = await router.read("a.pdf", make_pdf(["text", "text"]))
    assert pdf.ok and pdf.engines_used == ["marker"] and engines.calls == [("marker", "a.pdf", 2)]
    docx = await router.read("p.docx", make_docx())
    assert docx.ok and docx.source == "native" and "# Refund policy" in docx.markdown
    xlsx = await router.read("s.xlsx", make_xlsx())
    assert xlsx.ok and xlsx.source == "native"
    assert len(engines.calls) == 1  # Docling was never used, even though it is available


async def test_marker_mode_has_no_engine_behind_the_custom_reader(make_router):
    from docx import Document

    empty = io.BytesIO()
    Document().save(empty)
    engines = FakeEngines(order=("docling", "marker"), available=["docling", "marker"])
    router = make_router(engines, mode="marker")
    thin = await router.read("empty.docx", empty.getvalue())
    assert not thin.ok and engines.calls == []
    html = await router.read("page.html", b"<p>x</p>")
    assert not html.ok and engines.calls == []


async def test_marker_mode_reports_when_marker_is_not_available(make_router):
    result = await make_router(FakeEngines(order=("docling",), available=["docling"]), mode="marker").read("a.pdf", make_pdf(["text"]))
    assert not result.ok and "marker is not available" in result.reason


async def test_mode_can_be_chosen_per_request(make_router):
    engines = FakeEngines(order=("docling", "marker"), available=["docling", "marker"])
    router = make_router(engines)  # custom by default
    assert (await router.read("a.pdf", make_pdf(["text"]))).source == "native"
    assert (await router.read("a.pdf", make_pdf(["text"]), mode="docling")).engines_used == ["docling"]
    assert (await router.read("a.pdf", make_pdf(["text"]), mode="marker")).engines_used == ["marker"]


# -- office, text and other formats ----------------------------------------------------------------


async def test_docx_is_read_natively_in_order(make_router):
    engines = FakeEngines()
    result = await make_router(engines).read("p.docx", make_docx())
    text = result.markdown
    assert result.source == "native" and engines.calls == []
    assert text.index("# Refund policy") < text.index("Restocking fee") < text.index("Item | Fee") < text.index("Laptop | 10%")


async def test_pptx_and_xlsx_are_read_natively(make_router):
    router = make_router()
    pptx = (await router.read("d.pptx", make_pptx())).markdown
    assert "[Slide 1]" in pptx and "Ship connectors" in pptx and "Notes: Mention the pilot" in pptx and "[Slide 2]" in pptx
    xlsx = (await router.read("s.xlsx", make_xlsx())).markdown
    assert xlsx == "## Fees\nItem | Fee\nLaptop | 10%\nPhone | 5%"


async def test_a_thin_office_file_goes_to_the_engine(make_router):
    from docx import Document

    empty = io.BytesIO()
    Document().save(empty)
    engines = FakeEngines()
    result = await make_router(engines).read("empty.docx", empty.getvalue())
    assert result.source == "engine" and engines.calls[0][0] == "docling"


async def test_a_zip_bomb_is_refused_and_never_offered_to_an_engine(make_router):
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (20 * 1024 * 1024))
    engines = FakeEngines()
    result = await make_router(engines).read("bomb.docx", bomb.getvalue())
    assert not result.ok and "suspiciously" in result.reason and engines.calls == []


def test_zip_checks_refuse_other_compression_and_accept_normal_files():
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, "w", zipfile.ZIP_BZIP2) as archive:
        archive.writestr("a.xml", b"data")
    with pytest.raises(ReadError, match="compression"):
        check_zip_limits(packed.getvalue())
    check_zip_limits(make_docx())


async def test_plain_text_and_other_formats(make_router):
    engines = FakeEngines()
    router = make_router(engines)
    assert (await router.read("n.md", b"# Title\nbody")).markdown == "# Title\nbody"
    page = await router.read("page.html", b"<html><body>hello</body></html>")
    assert page.source == "engine" and engines.calls == [("docling", "page.html", 0)]
    none = await make_router(FakeEngines(order=())).read("page.html", b"<p>x</p>")
    assert not none.ok and "conversion engine is not available" in none.reason


# -- the worker pool -------------------------------------------------------------------------------


async def test_a_hung_worker_times_out_and_the_pool_recovers():
    pool = LightPool(max_workers=1, timeout=30)
    try:
        with pytest.raises(ReadError, match="too long"):
            await pool.run(sleep_then_return, 5, timeout=0.3)
        assert await pool.run(sleep_then_return, 0) == "late"
    finally:
        pool.shutdown()


# -- endpoint --------------------------------------------------------------------------------------


async def test_read_endpoint_reports_how_each_page_was_read(make_router):
    app = FastAPI()
    app.state.reading_router = make_router(FakeEngines())
    app.include_router(reading_routes)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/v1/documents/read", files={"file": ("a.pdf", make_pdf(["text", "scan"]), "application/pdf")})
        empty = await client.post("/v1/documents/read", files={"file": ("a.pdf", b"", "application/pdf")})
        whole = await client.post("/v1/documents/read", files={"file": ("a.pdf", make_pdf(["text"]), "application/pdf")}, data={"pdf_mode": "whole"})
        direct = await client.post("/v1/documents/read", files={"file": ("a.pdf", make_pdf(["text"]), "application/pdf")}, data={"mode": "docling"})
    body = response.json()
    assert body["ok"] is True and body["source"] == "mixed" and body["engines_used"] == ["docling"]
    assert [(p["page"], p["source"], p["reason"]) for p in body["pages"]] == [(1, "native", ""), (2, "engine", "image")]
    assert empty.status_code == 400
    assert whole.json()["source"] == "engine"
    assert direct.json()["engines_used"] == ["docling"]
